"""Deck state + python-pptx wrapper.

Decks are identified by a short `deck_id` (filename stem). State lives at
<root>/<deck_id>/deck.pptx; renders land at <root>/<deck_id>/renders/.
Open decks are kept in memory by DeckManager.

Element references: every shape has a stable `shape_id` int from python-pptx.
Tools take (deck_id, slide_index, element_id) to mutate.
"""
from __future__ import annotations
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor

from .units import inches, parse_color, emu_to_inches

SHAPE_TYPES = {
    "rectangle": MSO_SHAPE.RECTANGLE,
    "rounded_rectangle": MSO_SHAPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
    "oval": MSO_SHAPE.OVAL,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "right_triangle": MSO_SHAPE.RIGHT_TRIANGLE,
    "diamond": MSO_SHAPE.DIAMOND,
    "pentagon": MSO_SHAPE.REGULAR_PENTAGON,
    "hexagon": MSO_SHAPE.HEXAGON,
    "star": MSO_SHAPE.STAR_5_POINT,
    "arrow_right": MSO_SHAPE.RIGHT_ARROW,
    "arrow_left": MSO_SHAPE.LEFT_ARROW,
    "arrow_up": MSO_SHAPE.UP_ARROW,
    "arrow_down": MSO_SHAPE.DOWN_ARROW,
    "line": MSO_SHAPE.LINE_INVERSE,  # use add_image/connector for actual lines
    "callout": MSO_SHAPE.RECTANGULAR_CALLOUT,
    "cloud": MSO_SHAPE.CLOUD,
    "heart": MSO_SHAPE.HEART,
}

ALIGN_H = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY,
}

ALIGN_V = {
    "top": MSO_ANCHOR.TOP,
    "middle": MSO_ANCHOR.MIDDLE,
    "bottom": MSO_ANCHOR.BOTTOM,
}

_ID_RE = re.compile(r"^[a-zA-Z0-9_-]+$")


def _validate_id(deck_id: str) -> str:
    if not _ID_RE.match(deck_id):
        raise ValueError(f"deck_id must match [a-zA-Z0-9_-]+, got {deck_id!r}")
    return deck_id


@dataclass
class DeckHandle:
    deck_id: str
    path: Path
    prs: Presentation

    @property
    def render_dir(self) -> Path:
        return self.path.parent / "renders"


class DeckManager:
    def __init__(self, root: Path):
        self.root = Path(root).resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        self._open: dict[str, DeckHandle] = {}

    # -------- lifecycle --------
    def create(self, deck_id: str, width_in: float = 13.333, height_in: float = 7.5) -> DeckHandle:
        _validate_id(deck_id)
        deck_dir = self.root / deck_id
        if deck_dir.exists():
            raise FileExistsError(f"deck '{deck_id}' already exists at {deck_dir}")
        deck_dir.mkdir(parents=True)
        prs = Presentation()
        prs.slide_width = inches(width_in)
        prs.slide_height = inches(height_in)
        path = deck_dir / "deck.pptx"
        prs.save(str(path))
        h = DeckHandle(deck_id, path, prs)
        self._open[deck_id] = h
        return h

    def open(self, deck_id: str) -> DeckHandle:
        _validate_id(deck_id)
        if deck_id in self._open:
            return self._open[deck_id]
        path = self.root / deck_id / "deck.pptx"
        if not path.exists():
            raise FileNotFoundError(f"no deck '{deck_id}' at {path}")
        h = DeckHandle(deck_id, path, Presentation(str(path)))
        self._open[deck_id] = h
        return h

    def get(self, deck_id: str) -> DeckHandle:
        if deck_id in self._open:
            return self._open[deck_id]
        return self.open(deck_id)

    def close(self, deck_id: str) -> None:
        self._open.pop(deck_id, None)

    def save(self, deck_id: str) -> str:
        h = self.get(deck_id)
        h.prs.save(str(h.path))
        return str(h.path)

    def list_decks(self) -> list[dict]:
        out = []
        for d in sorted(self.root.iterdir()):
            p = d / "deck.pptx"
            if p.exists():
                out.append({"deck_id": d.name, "path": str(p), "open": d.name in self._open})
        return out

    # -------- slides --------
    def add_slide(self, deck_id: str, layout: int = 6) -> int:
        """layout 6 == blank in the default master. Returns 0-based slide index."""
        h = self.get(deck_id)
        layouts = h.prs.slide_layouts
        idx = max(0, min(layout, len(layouts) - 1))
        h.prs.slides.add_slide(layouts[idx])
        return len(h.prs.slides) - 1

    def delete_slide(self, deck_id: str, slide_index: int) -> None:
        h = self.get(deck_id)
        # python-pptx has no public delete; manipulate XML
        slides = h.prs.slides
        slide_id_lst = slides._sldIdLst  # noqa
        ids = list(slide_id_lst)
        if not (0 <= slide_index < len(ids)):
            raise IndexError(f"slide_index {slide_index} out of range")
        rId = ids[slide_index].rId
        h.prs.part.drop_rel(rId)
        slide_id_lst.remove(ids[slide_index])

    def reorder_slide(self, deck_id: str, from_index: int, to_index: int) -> None:
        h = self.get(deck_id)
        slide_id_lst = h.prs.slides._sldIdLst  # noqa
        ids = list(slide_id_lst)
        if not (0 <= from_index < len(ids)) or not (0 <= to_index < len(ids)):
            raise IndexError("slide index out of range")
        node = ids[from_index]
        slide_id_lst.remove(node)
        slide_id_lst.insert(to_index, node)

    def set_slide_background(self, deck_id: str, slide_index: int, color=None, image_path: str | None = None) -> None:
        h = self.get(deck_id)
        slide = h.prs.slides[slide_index]
        if color is not None:
            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = parse_color(color)
        if image_path is not None:
            # Insert as full-bleed picture, sent to back
            pic = slide.shapes.add_picture(
                image_path, 0, 0, h.prs.slide_width, h.prs.slide_height
            )
            spTree = slide.shapes._spTree  # noqa
            spTree.remove(pic._element)
            spTree.insert(2, pic._element)  # near the back

    # -------- elements --------
    def _slide(self, deck_id: str, slide_index: int):
        h = self.get(deck_id)
        if not (0 <= slide_index < len(h.prs.slides)):
            raise IndexError(f"slide_index {slide_index} out of range")
        return h.prs.slides[slide_index]

    def _find(self, deck_id: str, slide_index: int, element_id: int):
        slide = self._slide(deck_id, slide_index)
        for shape in slide.shapes:
            if int(shape.shape_id) == int(element_id):
                return shape
        raise KeyError(f"no element {element_id} on slide {slide_index}")

    def add_text_box(self, deck_id: str, slide_index: int, x: float, y: float,
                     w: float, h: float, text: str = "",
                     font: str | None = None, size_pt: float | None = None,
                     bold: bool | None = None, italic: bool | None = None,
                     color=None, align: str | None = None, valign: str | None = None,
                     auto_size: bool = False) -> int:
        slide = self._slide(deck_id, slide_index)
        tb = slide.shapes.add_textbox(inches(x), inches(y), inches(w), inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        if auto_size:
            from pptx.enum.text import MSO_AUTO_SIZE
            tf.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
        tf.text = text
        p = tf.paragraphs[0]
        if align in ALIGN_H:
            p.alignment = ALIGN_H[align]
        if valign in ALIGN_V:
            tf.vertical_anchor = ALIGN_V[valign]
        for run in p.runs:
            if font:
                run.font.name = font
            if size_pt:
                run.font.size = Pt(size_pt)
            if bold is not None:
                run.font.bold = bold
            if italic is not None:
                run.font.italic = italic
            if color is not None:
                run.font.color.rgb = parse_color(color)
        return int(tb.shape_id)

    def add_image(self, deck_id: str, slide_index: int, image_path: str,
                  x: float, y: float, w: float | None = None, h: float | None = None) -> int:
        slide = self._slide(deck_id, slide_index)
        kw = {}
        if w is not None:
            kw["width"] = inches(w)
        if h is not None:
            kw["height"] = inches(h)
        pic = slide.shapes.add_picture(image_path, inches(x), inches(y), **kw)
        return int(pic.shape_id)

    def add_shape(self, deck_id: str, slide_index: int, shape_type: str,
                  x: float, y: float, w: float, h: float,
                  fill=None, line_color=None, line_width_pt: float | None = None,
                  text: str | None = None) -> int:
        slide = self._slide(deck_id, slide_index)
        if shape_type not in SHAPE_TYPES:
            raise ValueError(f"unknown shape_type {shape_type!r}; one of {sorted(SHAPE_TYPES)}")
        sh = slide.shapes.add_shape(SHAPE_TYPES[shape_type], inches(x), inches(y), inches(w), inches(h))
        if fill is not None:
            sh.fill.solid()
            sh.fill.fore_color.rgb = parse_color(fill)
        if line_color is not None:
            sh.line.color.rgb = parse_color(line_color)
        if line_width_pt is not None:
            sh.line.width = Pt(line_width_pt)
        if text is not None:
            sh.text_frame.text = text
        return int(sh.shape_id)

    def set_text(self, deck_id: str, slide_index: int, element_id: int,
                 text: str | None = None, font: str | None = None,
                 size_pt: float | None = None, bold: bool | None = None,
                 italic: bool | None = None, underline: bool | None = None,
                 color=None, align: str | None = None, valign: str | None = None) -> None:
        sh = self._find(deck_id, slide_index, element_id)
        if not sh.has_text_frame:
            raise ValueError(f"element {element_id} has no text frame")
        tf = sh.text_frame
        if text is not None:
            tf.text = text
        if valign in ALIGN_V:
            tf.vertical_anchor = ALIGN_V[valign]
        for p in tf.paragraphs:
            if align in ALIGN_H:
                p.alignment = ALIGN_H[align]
            for run in p.runs:
                if font:
                    run.font.name = font
                if size_pt:
                    run.font.size = Pt(size_pt)
                if bold is not None:
                    run.font.bold = bold
                if italic is not None:
                    run.font.italic = italic
                if underline is not None:
                    run.font.underline = underline
                if color is not None:
                    run.font.color.rgb = parse_color(color)

    def set_position(self, deck_id: str, slide_index: int, element_id: int,
                     x: float | None = None, y: float | None = None) -> None:
        sh = self._find(deck_id, slide_index, element_id)
        if x is not None:
            sh.left = inches(x)
        if y is not None:
            sh.top = inches(y)

    def set_size(self, deck_id: str, slide_index: int, element_id: int,
                 w: float | None = None, h: float | None = None) -> None:
        sh = self._find(deck_id, slide_index, element_id)
        if w is not None:
            sh.width = inches(w)
        if h is not None:
            sh.height = inches(h)

    def set_rotation(self, deck_id: str, slide_index: int, element_id: int, degrees: float) -> None:
        sh = self._find(deck_id, slide_index, element_id)
        sh.rotation = float(degrees)

    def set_z_order(self, deck_id: str, slide_index: int, element_id: int, action: str) -> None:
        """action in {'front','back','forward','backward'}"""
        slide = self._slide(deck_id, slide_index)
        sh = self._find(deck_id, slide_index, element_id)
        spTree = slide.shapes._spTree  # noqa
        el = sh._element
        children = list(spTree)
        idx = children.index(el)
        spTree.remove(el)
        if action == "front":
            spTree.append(el)
        elif action == "back":
            spTree.insert(2, el)  # after nvGrpSpPr/grpSpPr
        elif action == "forward":
            spTree.insert(min(idx + 1, len(spTree)), el)
        elif action == "backward":
            spTree.insert(max(2, idx - 1), el)
        else:
            raise ValueError(f"action must be front|back|forward|backward, got {action!r}")

    def delete_element(self, deck_id: str, slide_index: int, element_id: int) -> None:
        slide = self._slide(deck_id, slide_index)
        sh = self._find(deck_id, slide_index, element_id)
        sh._element.getparent().remove(sh._element)

    # -------- inspection --------
    def get_deck_info(self, deck_id: str) -> dict:
        h = self.get(deck_id)
        return {
            "deck_id": deck_id,
            "path": str(h.path),
            "slide_count": len(h.prs.slides),
            "slide_size_inches": [emu_to_inches(h.prs.slide_width), emu_to_inches(h.prs.slide_height)],
        }

    def get_slide_info(self, deck_id: str, slide_index: int) -> dict:
        slide = self._slide(deck_id, slide_index)
        elements = []
        for sh in slide.shapes:
            info: dict[str, Any] = {
                "id": int(sh.shape_id),
                "name": sh.name,
                "type": str(sh.shape_type),
                "x_in": emu_to_inches(sh.left or 0),
                "y_in": emu_to_inches(sh.top or 0),
                "w_in": emu_to_inches(sh.width or 0),
                "h_in": emu_to_inches(sh.height or 0),
                "rotation": float(getattr(sh, "rotation", 0) or 0),
            }
            if sh.has_text_frame:
                info["text"] = sh.text_frame.text
            elements.append(info)
        return {"slide_index": slide_index, "elements": elements}

    def list_slides(self, deck_id: str) -> list[dict]:
        h = self.get(deck_id)
        return [{"index": i, "element_count": len(s.shapes)} for i, s in enumerate(h.prs.slides)]

    # -------- text styling inspection --------
    def get_text_style(self, deck_id: str, slide_index: int,
                      element_id: int | None = None,
                      search_text: str | None = None) -> dict:
        """Return text styling for a shape or for shapes matching search_text.

        Resolves placeholder inheritance from the slide layout, including the
        layout's prompt text (e.g. 'Presenter:') and its lstStyle defRPr —
        catches sizes/fonts/colors that python-pptx's run.font.* returns as None.
        """
        from xml.etree import ElementTree as ET
        A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

        def _emu_to_in(v):
            return None if v is None else round(int(v) / 914400, 4)

        def _runs_from_p(p_el):
            out = []
            for r in p_el.findall(A + "r"):
                rPr = r.find(A + "rPr")
                t_el = r.find(A + "t")
                info = {"text": (t_el.text or "") if t_el is not None else ""}
                if rPr is not None:
                    if rPr.get("sz"): info["size_pt"] = int(rPr.get("sz")) / 100
                    if rPr.get("b") is not None: info["bold"] = rPr.get("b") in ("1", "true")
                    if rPr.get("i") is not None: info["italic"] = rPr.get("i") in ("1", "true")
                    if rPr.get("u"): info["underline"] = rPr.get("u")
                    latin = rPr.find(A + "latin")
                    if latin is not None: info["font"] = latin.get("typeface")
                    solid = rPr.find(A + "solidFill")
                    if solid is not None:
                        srgb = solid.find(A + "srgbClr")
                        scheme = solid.find(A + "schemeClr")
                        if srgb is not None: info["color_rgb"] = srgb.get("val")
                        elif scheme is not None:
                            info["color_scheme"] = scheme.get("val")
                            lumMod = scheme.find(A + "lumMod")
                            lumOff = scheme.find(A + "lumOff")
                            if lumMod is not None: info["lumMod"] = int(lumMod.get("val")) / 1000
                            if lumOff is not None: info["lumOff"] = int(lumOff.get("val")) / 1000
                out.append(info)
            return out

        def _xfrm(sp_el):
            spPr = sp_el.find(P + "spPr")
            if spPr is None: return {}
            xfrm = spPr.find(A + "xfrm")
            if xfrm is None: return {}
            off = xfrm.find(A + "off"); ext = xfrm.find(A + "ext")
            o = {}
            if off is not None:
                o["x_in"] = _emu_to_in(off.get("x")); o["y_in"] = _emu_to_in(off.get("y"))
            if ext is not None:
                o["w_in"] = _emu_to_in(ext.get("cx")); o["h_in"] = _emu_to_in(ext.get("cy"))
            return o

        def _ph_idx(sp_el):
            ph = sp_el.find(f"{P}nvSpPr/{P}nvPr/{P}ph")
            return ph.get("idx") if ph is not None else None

        def _lst_default(sp_el):
            txBody = sp_el.find(P + "txBody")
            if txBody is None: return None
            lst = txBody.find(A + "lstStyle")
            if lst is None: return None
            lvl = lst.find(A + "lvl1pPr")
            if lvl is None: return None
            defR = lvl.find(A + "defRPr")
            if defR is None: return None
            out = {}
            if defR.get("sz"): out["size_pt"] = int(defR.get("sz")) / 100
            if defR.get("b") is not None: out["bold"] = defR.get("b") in ("1", "true")
            if defR.get("i") is not None: out["italic"] = defR.get("i") in ("1", "true")
            latin = defR.find(A + "latin")
            if latin is not None: out["font"] = latin.get("typeface")
            solid = defR.find(A + "solidFill")
            if solid is not None:
                srgb = solid.find(A + "srgbClr")
                scheme = solid.find(A + "schemeClr")
                if srgb is not None: out["color_rgb"] = srgb.get("val")
                elif scheme is not None:
                    out["color_scheme"] = scheme.get("val")
                    lumMod = scheme.find(A + "lumMod")
                    lumOff = scheme.find(A + "lumOff")
                    if lumMod is not None: out["lumMod"] = int(lumMod.get("val")) / 1000
                    if lumOff is not None: out["lumOff"] = int(lumOff.get("val")) / 1000
            return out

        slide = self._slide(deck_id, slide_index)
        slide_part = slide.part
        layout_part = slide_part.slide_layout.part if hasattr(slide_part, "slide_layout") else None
        try:
            layout_root = ET.fromstring(layout_part.blob)
        except Exception:
            layout_root = None
        slide_root = ET.fromstring(slide_part.blob)

        # Optionally filter shapes
        if element_id is not None:
            shapes = [s for s in slide.shapes if int(s.shape_id) == int(element_id)]
            if not shapes:
                raise KeyError(f"no element {element_id} on slide {slide_index}")
        else:
            shapes = list(slide.shapes)

        needle_lo = (search_text or "").lower()
        short = (search_text.split(":", 1)[-1].strip().lower() if search_text and ":" in search_text else needle_lo)
        prefix = (search_text.split(":", 1)[0].strip().lower() + ":") if search_text and ":" in search_text else None

        results = []
        for sh in shapes:
            if not sh.has_text_frame:
                continue
            text = sh.text_frame.text
            if search_text and needle_lo not in text.lower() and (not short or short not in text.lower()):
                continue
            # Find matching sp element in slide XML by id
            sp_el = None
            for sp in slide_root.iter(P + "sp"):
                cNvPr = sp.find(f"{P}nvSpPr/{P}cNvPr")
                if cNvPr is not None and cNvPr.get("id") == str(int(sh.shape_id)):
                    sp_el = sp; break
            paragraphs = []
            ph_idx = None
            if sp_el is not None:
                ph_idx = _ph_idx(sp_el)
                for p in sp_el.iter(A + "p"):
                    paragraphs.append(_runs_from_p(p))
            info = {
                "shape_id": int(sh.shape_id),
                "name": sh.name,
                "is_placeholder": sh.is_placeholder,
                "ph_idx": ph_idx,
                "x_in": _emu_to_in(sh.left), "y_in": _emu_to_in(sh.top),
                "w_in": _emu_to_in(sh.width), "h_in": _emu_to_in(sh.height),
                "text": text,
                "paragraphs": paragraphs,
            }
            # resolve inherited styling from layout placeholder of same idx
            if ph_idx is not None and layout_root is not None:
                for lsp in layout_root.iter(P + "sp"):
                    if _ph_idx(lsp) == ph_idx:
                        info["layout_ph"] = {
                            "text": "".join((t.text or "") for t in lsp.iter(A + "t")),
                            **_xfrm(lsp),
                            "paragraphs": [_runs_from_p(p) for p in lsp.iter(A + "p")],
                            "lstStyle_defRPr": _lst_default(lsp),
                        }
                        break
            results.append(info)

        # If a search prefix like "Presenter:" wasn't matched on any slide shape,
        # also surface matching layout-only labels (prompt text).
        layout_only = []
        if prefix and layout_root is not None and not any(prefix in r["text"].lower() for r in results):
            for lsp in layout_root.iter(P + "sp"):
                ltext = "".join((t.text or "") for t in lsp.iter(A + "t"))
                if prefix in ltext.lower():
                    layout_only.append({
                        "source": "layout_label",
                        "ph_idx": _ph_idx(lsp),
                        "text": ltext,
                        **_xfrm(lsp),
                        "paragraphs": [_runs_from_p(p) for p in lsp.iter(A + "p")],
                        "lstStyle_defRPr": _lst_default(lsp),
                    })

        return {"slide_index": slide_index, "matches": results, "layout_labels": layout_only}
